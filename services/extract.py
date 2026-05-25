"""Extract plain text from PDF files and web pages."""

from __future__ import annotations

import io
import json
import re
import time
from urllib.parse import urlparse

import httpx
from bs4 import BeautifulSoup, Comment
from pypdf import PdfReader


def _extract_pdf_pymupdf(data: bytes) -> str | None:
    """PyMuPDF often extracts more complete text and better reading order than pypdf alone."""
    try:
        import fitz
    except ImportError:
        return None
    if not data:
        return None
    try:
        doc = fitz.open(stream=data, filetype="pdf")
        try:
            chunks: list[str] = []
            for page in doc:
                try:
                    t = page.get_text(sort=True)
                except TypeError:
                    t = page.get_text()
                if t and t.strip():
                    chunks.append(t.strip())
            merged = "\n\n".join(chunks).strip()
            return merged if merged else None
        finally:
            doc.close()
    except Exception:
        return None


def _extract_pdf_pypdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        t = page.extract_text()
        if t:
            parts.append(t)
    return "\n\n".join(parts).strip()


def extract_pdf_bytes(data: bytes) -> str:
    """Try PyMuPDF first, then fall back to pypdf."""
    primary = _extract_pdf_pymupdf(data)
    if primary:
        return primary
    return _extract_pdf_pypdf(data)


def _decode_text_file(data: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "gbk", "gb2312"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _extract_docx_bytes(data: bytes) -> str:
    import io

    from docx import Document

    doc = Document(io.BytesIO(data))
    parts: list[str] = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts).strip()


def extract_document_bytes(filename: str, data: bytes) -> str:
    """PDF / Word(docx) / Markdown / 纯文本。"""
    lower = (filename or "").lower()
    if lower.endswith(".pdf"):
        return extract_pdf_bytes(data)
    if lower.endswith(".docx"):
        return _extract_docx_bytes(data)
    if lower.endswith(".doc"):
        raise ValueError("不支持旧版 .doc，请在 Word 中另存为 .docx 后上传")
    if lower.endswith((".md", ".markdown", ".txt", ".text")):
        return _decode_text_file(data).strip()
    raise ValueError(f"不支持的文件类型：{filename}（支持 PDF、Word(docx)、Markdown、TXT）")


def is_supported_document_filename(filename: str) -> bool:
    lower = (filename or "").lower()
    return lower.endswith(
        (".pdf", ".docx", ".md", ".markdown", ".txt", ".text"),
    )


_WS_RE = re.compile(r"[ \t]+")


def _norm_line(s: str) -> str:
    return _WS_RE.sub(" ", s.strip())


def _dedupe_lines(blocks: list[str], min_len: int = 2) -> str:
    """Merge sections and drop exact duplicate lines (keep first occurrence order)."""
    seen: set[str] = set()
    out: list[str] = []
    for block in blocks:
        for line in block.splitlines():
            n = _norm_line(line)
            if len(n) < min_len or n in seen:
                continue
            seen.add(n)
            out.append(n)
    return "\n".join(out).strip()


def _collect_jsonld_texts(obj: object, sink: list[str]) -> None:
    keys = frozenset(
        {
            "articleBody",
            "description",
            "headline",
            "name",
            "text",
            "abstract",
            "content",
        }
    )
    if isinstance(obj, dict):
        for k, v in obj.items():
            if isinstance(v, str) and k in keys and len(v.strip()) > 15:
                sink.append(v.strip())
            else:
                _collect_jsonld_texts(v, sink)
    elif isinstance(obj, list):
        for x in obj:
            _collect_jsonld_texts(x, sink)


_META_NAMES = frozenset(
    {
        "description",
        "keywords",
        "author",
        "news_keywords",
    }
)


def _meta_content_is_useful(tag, content: str) -> bool:
    """Skip viewport、缓存指令等非正文 meta。"""
    c = content.strip()
    if len(c) < 8:
        return False
    name = (tag.get("name") or "").strip().lower()
    prop = (tag.get("property") or "").strip().lower()
    if name in _META_NAMES:
        return True
    if prop.startswith(("og:", "twitter:", "article:")):
        return True
    if "title" in prop and prop.startswith("og:"):
        return True
    if name in ("robots", "viewport", "referrer") or "verification" in name:
        return False
    if "width=device-width" in c or c.lower() in ("no-cache", "0", "ie=edge"):
        return False
    if len(c) > 120 and not c.startswith("http"):
        return True
    return False


def _extract_meta_and_jsonld(soup: BeautifulSoup) -> list[str]:
    """
    Title + 文案类 meta（description / keywords / og / twitter 等）。
    SPA 壳站（如 moonshot.cn）静态 HTML 里正文在 JS，但 meta 里常有简介与关键词。
    """
    parts: list[str] = []
    seen: set[str] = set()

    def add(s: str, min_len: int = 1) -> None:
        t = s.strip()
        if len(t) < min_len or t in seen:
            return
        seen.add(t)
        parts.append(t)

    if soup.title and soup.title.string:
        add(soup.title.string.strip(), min_len=1)

    for tag in soup.find_all("meta"):
        c = tag.get("content")
        if not isinstance(c, str) or not c.strip():
            continue
        if not _meta_content_is_useful(tag, c):
            continue
        add(c.strip(), min_len=2)

    for script in soup.find_all("script", type=lambda x: x and "ld+json" in str(x).lower()):
        raw = (script.string or script.get_text() or "").strip()
        if not raw:
            continue
        try:
            data = json.loads(raw)
        except json.JSONDecodeError:
            continue
        ld_parts: list[str] = []
        _collect_jsonld_texts(data, ld_parts)
        for lp in ld_parts:
            add(lp, min_len=10)

    return parts


def _extract_trafilatura(html: str, url: str) -> str | None:
    try:
        import trafilatura
    except ImportError:
        return None
    try:
        t = trafilatura.extract(
            html,
            url=url,
            include_tables=True,
            include_comments=True,
            favor_recall=True,
            include_links=True,
        )
        return t.strip() if t and t.strip() else None
    except Exception:
        return None


def _style_attr_hides_element(style_val: str | None) -> bool:
    """display:none / visibility:hidden 等（用于判断隐藏域，不据此一律删除）。"""
    if not style_val:
        return False
    s = re.sub(r"\s+", "", (style_val or "").lower())
    return "display:none" in s or "visibility:hidden" in s


def _textarea_is_hidden(ta) -> bool:
    return ta.get("hidden") is not None or _style_attr_hides_element(ta.get("style"))


# 典型「整段样式表」片段，JSON/正文很少同时命中多项
_CSS_BUNDLE_MARKERS = (
    "box-sizing",
    "margin:",
    "padding:",
    "font-size:",
    "background-",
    "border:",
    "line-height:",
    "{",
)


def _raw_smells_like_css_bundle(raw: str) -> bool:
    """仅当很像压缩 CSS 时再丢，避免误伤隐藏域里的 JSON（如热搜数据）。"""
    if len(raw) < 400:
        return False
    low = raw.lower()
    hits = sum(1 for m in _CSS_BUNDLE_MARKERS if m in low)
    return hits >= 5 and raw.count("{") >= 6 and raw.count(";") >= 12


def _decompose_css_only_hidden_textareas(soup: BeautifulSoup) -> None:
    """
    只去掉「隐藏 + 明显是 CSS 包」的 textarea（如百度 s_is_result_css）。
    不再去掉所有 display:none 的 textarea，以免删掉 hotsearch_data 等有用 JSON。
    """
    for ta in soup.find_all("textarea"):
        if not _textarea_is_hidden(ta):
            continue
        tid = (ta.get("id") or "").lower()
        raw = (ta.get_text() or "").strip()
        if "_css" in tid or tid.endswith("stylesheet") or "style_text" in tid:
            ta.decompose()
            continue
        if _raw_smells_like_css_bundle(raw):
            ta.decompose()


def _extract_full_body_text(soup: BeautifulSoup) -> str:
    """After heavy tags removed: full visible text in document order."""
    for c in soup.find_all(string=lambda t: isinstance(t, Comment)):
        c.extract()
    _decompose_css_only_hidden_textareas(soup)
    for tag in soup(["script", "style", "noscript", "template"]):
        tag.decompose()
    root = soup.body or soup
    text = root.get_text("\n", strip=False)
    lines = [_norm_line(ln) for ln in text.splitlines()]
    return "\n".join(ln for ln in lines if ln).strip()


def _fetch_html_playwright(url: str, timeout: float) -> tuple[str, str]:
    """
    返回 (序列化 HTML, body 可见 inner_text)。SPA 站点用 inner_text 往往比单独解析 HTML 更完整。
    """
    try:
        from playwright.sync_api import sync_playwright
    except ImportError as e:
        raise RuntimeError(
            "已开启「渲染 JavaScript」，但当前运行 uvicorn 的 Python 里找不到 playwright 模块。\n\n"
            "你在终端里用「系统全局」pip / 家目录下 playwright 安装的，不会进项目的 .venv。\n"
            "请在项目目录（含 main.py 与 .venv 的文件夹）下执行：\n"
            "  .venv/bin/pip install playwright\n"
            "  .venv/bin/python -m playwright install chromium\n\n"
            "装完后重启 uvicorn。可在浏览器打开 /api/playwright-status 查看当前后端用的 Python 路径。"
        ) from e
    timeout_ms = int(max(15.0, timeout) * 1000)
    ua = (
        "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36"
    )
    body_inner = ""
    with sync_playwright() as p:
        try:
            browser = p.chromium.launch(
                headless=True,
                args=[
                    "--disable-blink-features=AutomationControlled",
                    "--disable-dev-shm-usage",
                    "--no-sandbox",
                ],
            )
        except Exception as e:
            raise RuntimeError(
                "Playwright 模块已加载，但无法启动 Chromium（多半未在「当前 Python」下下载浏览器）。\n"
                "请在项目目录执行：\n"
                "  .venv/bin/python -m playwright install chromium\n\n"
                f"原始错误：{e}"
            ) from e
        try:
            context = browser.new_context(
                user_agent=ua,
                viewport={"width": 1365, "height": 900},
                locale="zh-CN",
                extra_http_headers={
                    "Accept-Language": "zh-CN,zh;q=0.9,en-US;q=0.8,en;q=0.7",
                },
            )
            page = context.new_page()
            page.add_init_script(
                "Object.defineProperty(navigator, 'webdriver', {get: () => undefined});"
            )
            page.goto(url, wait_until="domcontentloaded", timeout=timeout_ms)
            try:
                page.wait_for_function(
                    """() => {
                      const el = document.querySelector('#root, main, [role=main], #app, body');
                      if (!el) return false;
                      const t = (el.innerText || '').trim();
                      return t.length > 80;
                    }""",
                    timeout=min(28000, max(8000, timeout_ms - 12000)),
                )
            except Exception:
                pass
            time.sleep(1.2)
            try:
                page.wait_for_load_state("networkidle", timeout=15000)
            except Exception:
                pass
            time.sleep(0.8)
            try:
                body_inner = page.inner_text("body", timeout=8000)
            except Exception:
                try:
                    body_inner = page.evaluate(
                        "() => (document.body && document.body.innerText) ? document.body.innerText : ''"
                    )
                except Exception:
                    body_inner = ""
            html = page.content()
            context.close()
        finally:
            browser.close()
    return html, (body_inner or "").strip()


def _html_to_text_blocks(
    html: str,
    url: str,
    *,
    playwright_body_text: str | None = None,
) -> str:
    soup = BeautifulSoup(html, "lxml")
    meta_blocks = _extract_meta_and_jsonld(soup)
    traf = _extract_trafilatura(html, url)
    soup2 = BeautifulSoup(html, "lxml")
    body_full = _extract_full_body_text(soup2)

    sections: list[str] = []
    if meta_blocks:
        sections.append("\n".join(meta_blocks))
    if traf:
        sections.append(traf)
    if body_full:
        sections.append(body_full)
    if playwright_body_text and playwright_body_text.strip():
        sections.append(playwright_body_text.strip())

    if not sections:
        return ""
    return _dedupe_lines(sections)


def extract_web_url(url: str, timeout: float = 45.0, render_js: bool = False) -> str:
    """
    尽量抓取网页文本。SPA（如 moonshot.cn）静态 HTML 里往往只有 <title> 与 meta，正文在 JS 里。

    - render_js=False：HTTP 拉取 + 文案类 meta + trafilatura + 整页可见文本（适合文档站、新闻站）。
    - render_js=True：用 Chromium 执行页面后再同样抽取（需安装 playwright + playwright install chromium）。
    """
    parsed = urlparse(url.strip())
    if parsed.scheme not in ("http", "https"):
        raise ValueError("仅支持 http/https 链接")

    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
            "(KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36"
        ),
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
        "Cache-Control": "no-cache",
    }

    if render_js:
        html, visible = _fetch_html_playwright(url, timeout)
        return _html_to_text_blocks(html, url, playwright_body_text=visible)
    with httpx.Client(follow_redirects=True, timeout=timeout) as client:
        r = client.get(url, headers=headers)
        r.raise_for_status()
        html = r.text

    return _html_to_text_blocks(html, url)
